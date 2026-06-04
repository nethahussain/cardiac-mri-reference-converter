#!/usr/bin/env python3
"""Generate the English documentation (.docx) for the Cardiac MRI Reference Converter:
   docs/Calculations_Documentation.docx  and  docs/Reference_Values.docx
The reference tables are built directly from the validated cmr_reference_ranges.json.
"""
import json, os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

os.makedirs('docs', exist_ok=True)
data = json.load(open('cmr_reference_ranges.json'))
REFS = data['REFS']

ETHNICITIES = ["White", "Black", "South Asian", "Chinese", "Mixed/Other"]
SEXES = [("F", "Female"), ("M", "Male")]
SEGS = [("smooth", "Smooth"), ("papillary", "Papillary"), ("anatomical", "Anatomical")]
IDXS = [("bsa", "body surface area"), ("height", "height")]
AGES = ["18-29", "30-39", "40-49", "50-59", "60-69", "70+"]
CHAMBERS = [
    ("Left ventricle", "LV", ["LVEDV", "LVEDVi", "LVESV", "LVESVi", "LVSV", "LVSVi",
                              "LVCO", "LVEF", "LVM diast", "LVMi diast", "LVM syst", "LVMi syst"]),
    ("Right ventricle", "RV", ["RVEDV", "RVEDVi", "RVESV", "RVESVi", "RVSV", "RVSVi", "RVCO", "RVEF"]),
    ("Left atrium", "LA", ["LAESV", "LAESVi", "LA max", "LA max i", "LAEF"]),
    ("Right atrium", "RA", ["RAESV", "RAESVi", "RA max", "RA max i", "RAEF"]),
]

ABS_VOL = {"LVEDV", "LVESV", "LVSV", "RVEDV", "RVESV", "RVSV", "LAESV", "LA max", "RAESV", "RA max"}
IDX_VOL = {"LVEDVi", "LVESVi", "LVSVi", "RVEDVi", "RVESVi", "RVSVi", "LAESVi", "LA max i", "RAESVi", "RA max i"}
ABS_MASS = {"LVM diast", "LVM syst"}
IDX_MASS = {"LVMi diast", "LVMi syst"}
CO = {"LVCO", "RVCO"}
EF = {"LVEF", "RVEF", "LAEF", "RAEF"}

DESC = {
    "LVEDV": "LV end-diastolic volume", "LVEDVi": "LV end-diastolic volume index",
    "LVESV": "LV end-systolic volume", "LVESVi": "LV end-systolic volume index",
    "LVSV": "LV stroke volume", "LVSVi": "LV stroke volume index",
    "LVCO": "LV cardiac output", "LVEF": "LV ejection fraction",
    "LVM diast": "LV myocardial mass, end-diastole", "LVMi diast": "LV myocardial mass index, end-diastole",
    "LVM syst": "LV myocardial mass, end-systole", "LVMi syst": "LV myocardial mass index, end-systole",
    "RVEDV": "RV end-diastolic volume", "RVEDVi": "RV end-diastolic volume index",
    "RVESV": "RV end-systolic volume", "RVESVi": "RV end-systolic volume index",
    "RVSV": "RV stroke volume", "RVSVi": "RV stroke volume index",
    "RVCO": "RV cardiac output", "RVEF": "RV ejection fraction",
    "LAESV": "LA end-systolic (maximal) volume", "LAESVi": "LA end-systolic volume index",
    "LA max": "Maximal LA volume", "LA max i": "Maximal LA volume index", "LAEF": "LA ejection fraction",
    "RAESV": "RA end-systolic (maximal) volume", "RAESVi": "RA end-systolic volume index",
    "RA max": "Maximal RA volume", "RA max i": "Maximal RA volume index", "RAEF": "RA ejection fraction",
}

def unit(key, idx):
    if key in ABS_VOL: return "ml"
    if key in ABS_MASS: return "g"
    if key in CO: return "l/min"
    if key in EF: return "%"
    if key in IDX_VOL: return "ml/m²" if idx == "bsa" else "ml/m"
    if key in IDX_MASS: return "g/m²" if idx == "bsa" else "g/m"
    return ""

# ---------------------------------------------------------------- Calculations doc
def build_calculations():
    d = Document()
    d.add_heading("Cardiac MRI Reference Converter — Calculations", 0)
    d.add_paragraph("Documentation of every calculation performed by the tool. Reference "
                    "ranges are the 95% prediction intervals from Raisi-Estabragh et al. "
                    "(2024), Healthy Hearts Consortium (JACC Cardiovasc Imaging; "
                    "doi:10.1016/j.jcmg.2024.01.009).")

    d.add_heading("1. Body surface area (BSA) — Mosteller formula", 1)
    d.add_paragraph("BSA = √(weight × height / 3600), with weight in kg, height in cm, "
                    "and the result in m².")
    d.add_paragraph("Example: 73 kg, 180 cm → √(73 × 180 / 3600) = √3.65 = 1.91 m².")

    d.add_heading("2. Age and age group", 1)
    d.add_paragraph("Age is the difference between the current date and the date of birth "
                    "(reduced by 1 if this year's birthday has not yet occurred). Age is then "
                    "mapped to a reference age band:")
    t = d.add_table(rows=1, cols=2); t.style = "Table Grid"
    t.rows[0].cells[0].text = "Age"; t.rows[0].cells[1].text = "Age group"
    for a, g in [("18–29", "18-29"), ("30–39", "30-39"), ("40–49", "40-49"),
                 ("50–59", "50-59"), ("60–69", "60-69"), ("70+", "70+")]:
        r = t.add_row().cells; r[0].text = a; r[1].text = g
    d.add_paragraph("Below 18 years: no reference group available. Some ethnicity groups also "
                    "lack data for the younger bands; in those cases values are shown without a "
                    "reference comparison.")

    d.add_heading("3. Stroke volume (SV)", 1)
    d.add_paragraph("SV = EDV − ESV, computed separately for the left (LV) and right (RV) "
                    "ventricle. EDV = end-diastolic volume, ESV = end-systolic volume.")

    d.add_heading("4. Ejection fraction (EF)", 1)
    d.add_paragraph("EF = (SV / EDV) × 100 = ((EDV − ESV) / EDV) × 100, in percent — "
                    "the fraction of blood pumped out per beat.")

    d.add_heading("5. Cardiac output (CO)", 1)
    d.add_paragraph("CO = SV × heart rate / 1000, in l/min (SV in ml, heart rate in beats/min). "
                    "Cardiac output is only shown when a heart rate is provided.")

    d.add_heading("6. Body-size indexing", 1)
    d.add_paragraph("Each volume and mass is normalised to body size. The tool offers two "
                    "indexing methods, matching the two body-size corrections in the publication:")
    t = d.add_table(rows=1, cols=3); t.style = "Table Grid"
    for i, h in enumerate(["Method", "Formula", "Units"]): t.rows[0].cells[i].text = h
    for m, f, u in [("Body surface area", "value ÷ BSA", "ml/m², g/m²"),
                    ("Height", "value ÷ height (m)", "ml/m, g/m")]:
        r = t.add_row().cells; r[0].text = m; r[1].text = f; r[2].text = u

    d.add_heading("7. Atrial ejection fraction", 1)
    d.add_paragraph("LAEF and RAEF = (V_max − V_min) / V_max × 100, using the maximal and "
                    "minimal atrial volumes, in percent.")

    d.add_heading("8. Segmentation methods", 1)
    d.add_paragraph("The publication provides reference ranges for three predefined segmentation "
                    "protocols. They differ chiefly in what is included in the myocardial mass "
                    "(and, to a lesser extent, in the cavity volumes). Select the method that "
                    "matches how your measurements were obtained.")
    t = d.add_table(rows=1, cols=2); t.style = "Table Grid"
    t.rows[0].cells[0].text = "Method"; t.rows[0].cells[1].text = "Definition"
    for m, desc in [("Smooth", "Smooth endocardial contours; papillary muscles and trabeculae "
                                "excluded from the LV mass."),
                    ("Papillary", "Smooth endocardial contours with papillary muscles (but not "
                                  "trabeculae) included in the LV mass."),
                    ("Anatomical", "Both papillary muscles and trabeculae included in the LV mass.")]:
        r = t.add_row().cells; r[0].text = m; r[1].text = desc
    d.add_paragraph("Mass therefore increases from smooth → papillary → anatomical, while the "
                    "corresponding cavity volumes decrease slightly.")

    d.add_heading("9. Reference comparison", 1)
    d.add_paragraph("Each computed value is compared with the [min, max] range for the patient's "
                    "ethnicity, sex, age group, segmentation method and indexing:")
    t = d.add_table(rows=1, cols=3); t.style = "Table Grid"
    for i, h in enumerate(["Condition", "Status", "Flag"]): t.rows[0].cells[i].text = h
    for c, s, fl in [("value < min", "Low", "* (asterisk)"),
                     ("min ≤ value ≤ max", "Normal", "none"),
                     ("value > max", "High", "* (asterisk)")]:
        r = t.add_row().cells; r[0].text = c; r[1].text = s; r[2].text = fl

    d.add_heading("10. Visual gauge", 1)
    d.add_paragraph("The gauge maps a value to a percentage position: padding = (max − min) × 0.35; "
                    "gauge_min = min − padding; gauge_max = max + padding; "
                    "position = ((value − gauge_min) / (gauge_max − gauge_min)) × 100%, clamped to "
                    "0–100%. The green band marks the normal range.")

    d.add_heading("Workflow summary", 1)
    for step in ["Patient data is entered (sex, ethnicity, date of birth, weight, height, heart rate).",
                 "BSA is calculated with the Mosteller formula.",
                 "Segmentation method and indexing are selected.",
                 "Measurements are entered for each chamber, or imported from XML.",
                 "SV, EF and CO are derived; atrial EF is derived from max and min volumes.",
                 "All volumes and masses are indexed by BSA or height.",
                 "Reference ranges are looked up by ethnicity, sex, age group, segmentation and indexing.",
                 "Out-of-range values are flagged, and a report is generated for copying."]:
        d.add_paragraph(step, style="List Bullet")

    out = "docs/Calculations_Documentation.docx"; d.save(out); return out

# ---------------------------------------------------------------- Reference values doc
def build_reference_values():
    d = Document()
    # compact default font
    st = d.styles["Normal"]; st.font.size = Pt(9)
    d.add_heading("Cardiac MRI Reference Converter — Reference Values", 0)
    d.add_paragraph("Reference ranges (95% prediction intervals; each cell shows min–max) from:")
    p = d.add_paragraph(); r = p.add_run("Raisi-Estabragh Z, et al. Cardiovascular Magnetic "
        "Resonance Reference Ranges From the Healthy Hearts Consortium. JACC Cardiovasc Imaging. "
        "2024. doi:10.1016/j.jcmg.2024.01.009"); r.italic = True
    d.add_paragraph("Reference values released under CC0 1.0. Tables are organised by segmentation "
                    "method, then by body-size indexing, then by ethnicity and sex. Some ethnicity "
                    "groups lack data for the younger age bands.")

    d.add_heading("Parameter glossary", 1)
    g = d.add_table(rows=1, cols=4); g.style = "Table Grid"
    for i, h in enumerate(["Parameter", "Description", "Unit (BSA)", "Unit (height)"]):
        g.rows[0].cells[i].text = h
    seen = []
    for _, _, keys in CHAMBERS:
        for k in keys:
            if k in seen: continue
            seen.append(k)
            row = g.add_row().cells
            row[0].text = k; row[1].text = DESC.get(k, ""); row[2].text = unit(k, "bsa"); row[3].text = unit(k, "height")

    for seg_key, seg_name in SEGS:
        d.add_page_break()
        d.add_heading(f"{seg_name} segmentation", 1)
        for idx_key, idx_name in IDXS:
            d.add_heading(f"Indexed by {idx_name}", 2)
            for eth in ETHNICITIES:
                for sex_key, sex_name in SEXES:
                    try:
                        node = REFS[eth][sex_key][seg_key][idx_key]
                    except KeyError:
                        continue
                    ages = [a for a in AGES if any(node.get(k, {}).get(a) for _, _, ks in CHAMBERS for k in ks)]
                    if not ages:
                        continue
                    d.add_heading(f"{eth} — {sex_name}", 3)
                    tbl = d.add_table(rows=1, cols=1 + len(ages)); tbl.style = "Table Grid"
                    hdr = tbl.rows[0].cells
                    hdr[0].text = "Parameter"
                    for i, a in enumerate(ages): hdr[1 + i].text = a
                    for ch_name, ch_abbr, keys in CHAMBERS:
                        # chamber sub-header row (merged)
                        cr = tbl.add_row().cells
                        merged = cr[0]
                        for c in cr[1:]:
                            merged = merged.merge(c)
                        merged.text = ch_name
                        for para in merged.paragraphs:
                            for run in para.runs: run.bold = True
                        for k in keys:
                            if k not in node:
                                continue
                            rc = tbl.add_row().cells
                            u = unit(k, idx_key)
                            rc[0].text = f"{k} ({u})" if u else k
                            for i, a in enumerate(ages):
                                v = node[k].get(a)
                                rc[1 + i].text = f"{v[0]}–{v[1]}" if v else "—"
    out = "docs/Reference_Values.docx"; d.save(out); return out

c = build_calculations()
r = build_reference_values()
print("Wrote", c, os.path.getsize(c), "bytes")
print("Wrote", r, os.path.getsize(r), "bytes")
